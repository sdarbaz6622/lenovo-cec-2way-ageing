from html import escape
from io import BytesIO

import numpy as np
import openpyxl
import pandas as pd
import streamlit as st

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# STREAMLIT PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Lenovo CEC 2-Way Email Ageing Generator",
    page_icon="📧",
    layout="wide",
)


# ============================================================
# CUSTOM UI
# ============================================================

st.markdown(
    """
    <style>

    .main-title {
        font-size: 28px;
        font-weight: 700;
        color: #1F4E78;
        margin-bottom: 4px;
    }

    .sub-text {
        font-size: 14px;
        color: #666666;
        margin-bottom: 20px;
    }

    .stDownloadButton button {
        background-color: #1F4E78;
        color: white;
        font-weight: 700;
        border-radius: 6px;
        border: none;
    }

    .stDownloadButton button:hover {
        background-color: #163A5C;
        color: white;
    }

    .metric-box {
        padding: 14px;
        border-radius: 8px;
        background-color: #F3F7FA;
        border-left: 5px solid #1F4E78;
        margin-bottom: 10px;
    }

    .metric-title {
        font-size: 12px;
        color: #666666;
    }

    .metric-value {
        font-size: 24px;
        font-weight: 700;
        color: #1F4E78;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


st.markdown(
    '<div class="main-title">Lenovo CEC - 2-Way Email Ageing Report Generator</div>',
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="sub-text">
    Upload the latest raw Excel file to generate the TL-wise summary,
    TL & Agent ageing report, formatted Excel workbook and
    Outlook-friendly email.
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# HEADCOUNT / TEAM LEAD MAPPING
# ============================================================

# IMPORTANT:
# Keep your complete existing hc_dict here.
# I have retained the mapping you supplied.

hc_dict = {

    "Alex Xavier": "Alex Xavier",
    "Sumith Kumar Dey": "Sreejith Menon (Support Function)",
    "SHIVAKUMAR SHIRSHI": "A Afrid",
    "Amog K": "Sreejith Menon (Support Function)",
    "M Mohammed Khaja": "Sreejith Menon (Support Function)",
    "Illahin Naaz": "A Afrid",
    "Nikitha Kumari": "A Afrid",
    "Jisharaj K": "Sreejith Menon (Support Function)",
    "Aamir Hassan": "A Afrid",
    "Varshikey Tiwari": "A Afrid",
    "Mohan Kumar": "A Afrid",
    "Rohith N Shetty": "A Afrid",
    "Madhusudan Singh": "A Afrid",

    "Priyanshu kumar Singh": "Sreejith Menon",

    "Aman Kumar": "Augustine Joseph",
    "Ashish Kashyap": "Augustine Joseph",
    "Augustine Joseph": "Augustine Joseph",
    "Bijit Sen": "Augustine Joseph",
    "Harshith Singh N R": "Augustine Joseph",
    "Prakash Ambli": "Augustine Joseph",
    "Shaik Imran": "Augustine Joseph",
    "Shaik Mubarak": "Augustine Joseph",
    "Snehan Thejaswi T": "Augustine Joseph",

    "Devapatla Swetha": "Avrel Fernandez",
    "Ashish Kumar": "Avrel Fernandez",
    "Avrel Fernandez": "Avrel Fernandez",
    "Chandan M B": "Avrel Fernandez",
    "Chandana S": "Avrel Fernandez",
    "Mdshamshad Iqbal": "Avrel Fernandez",
    "Nivedita Kamate": "Avrel Fernandez",
    "Suresh Panda": "Avrel Fernandez",

    "Dinesh S": "Deepak Ghosh (Workstation)",
    "Chandra Prakash Singh": "Deepak Ghosh (Workstation)",
    "Deepak Ghosh": "Deepak Ghosh (Workstation)",
    "Hitesh Diliprao Deshmukh": "Deepak Ghosh (Workstation)",
    "Syed Hussain": "Deepak Ghosh (Workstation)",
    "Subham Prasad": "Deepak Ghosh (Workstation)",
    "Subrata Jana": "Deepak Ghosh (Workstation)",
    "Vinay Kumar Sharma": "Deepak Ghosh (Workstation)",

    "Anand Divatai": "Ram Mohan Singh",
    "Bijay Chettri": "Ram Mohan Singh",
    "Husna C": "Ram Mohan Singh",
    "Jyoti Nandini": "Ram Mohan Singh",
    "Manjushree Manjushree": "Ram Mohan Singh",
    "Miyani Smit": "Ram Mohan Singh",
    "Ram Mohan Singh": "Ram Mohan Singh",
    "Roshan Kumar": "Ram Mohan Singh",
    "Sha Ghousia": "Ram Mohan Singh",

    "Ashish Seet": "Seema Lal",
    "Debasish Panda": "Seema Lal",
    "Debasmita Samal": "Seema Lal",
    "Geetika Garnaik": "Seema Lal",
    "Paramita Ghosh": "Seema Lal",
    "Romeo Fernandez": "Seema Lal",
    "Rushikesh Powar": "Seema Lal",
    "Sandeep Das": "Seema Lal",
    "Sarvesh Singh": "Seema Lal",
    "Seema Lal": "Seema Lal",
    "Shafiulla Sameer": "Seema Lal",
    "Velisetty Kumar": "Seema Lal",

    "Anees Basha": "Manikanta KR",
    "Arshit Choubey": "Manikanta KR",
    "Imad Ulla Khan": "Manikanta KR",
    "Kanire Ravindra": "Manikanta KR",
    "Muhammed Ameen Sheikh": "Manikanta KR",
    "Naveen R": "Manikanta KR",
    "Praveen Kumar": "Manikanta KR",
    "Puranik Tippanna": "Manikanta KR",
    "Shrikant S": "Manikanta KR",
    "Siva Venkatraman": "Manikanta KR",
    "Solomon D": "Manikanta KR",

    "Aliya H": "Sreejith Menon",
    "Preeti Barik": "Sreejith Menon",
    "Priyambada Mohapatra": "Sreejith Menon",
    "Ram Govind": "Sreejith Menon",
    "Richa Panwar": "Sreejith Menon",
    "Rohini Bahubali Munnoli": "Sreejith Menon",
    "Sreejith Menon": "Sreejith Menon",
    "Sudharani Sudharani": "Sreejith Menon",

    "CHAITY DEBNATH": "Syed Arbaz",
    "Mitali Basumatary": "Syed Arbaz",
    "Mohammed Abdul Razzaq": "Syed Arbaz",
    "Shaik Ruhina Arshiya": "Syed Arbaz",
    "Shifa Mohammedi": "Syed Arbaz",
    "Syed Arbaz": "Syed Arbaz",
    "Vaishnavi Vaman Kulkarni": "Syed Arbaz",

    "A r riwan Ahmed": "Vikas B Bhovi",
    "ABHISHEK KUMAR": "Vikas B Bhovi",
    "Govindsa Ankita": "Vikas B Bhovi",
    "Mohammed Saqib": "Vikas B Bhovi",
    "Saniya Saniya": "Vikas B Bhovi",
    "Sneha Sneha": "Vikas B Bhovi",
    "Sneha Patil": "Vikas B Bhovi",
    "Vikas B Bhovi": "Vikas B Bhovi",
    "Vikrant Kumar": "Vikas B Bhovi",
    "Vinay Kumar Joshi": "Vikas B Bhovi",

    "Akshay Anil Mardhekar": "Think (L1.5)",
    "Amol Shioshankar Bhongade": "Think (L1.5)",
    "Arul S": "Think (L1.5)",
    "Hemavathi S": "Think (L1.5)",
    "Jagatheesh A": "Think (L1.5)",
    "Kiran J S": "Think (L1.5)",
    "Liyakathali Mohameth Khan": "Think (L1.5)",
    "Marwan Fahad Faiz Jabbar": "Think (L1.5)",
    "Md Khaja": "Think (L1.5)",
    "Mohammad Sajid": "Think (L1.5)",
    "Narsim Raj T S": "Think (L1.5)",
    "Naveen shivram Naik": "Think (L1.5)",
    "Raju Mondal": "Think (L1.5)",
    "Rammilan Pandit": "Think (L1.5)",
    "Santosh Kumar Biredar": "Think (L1.5)",
    "Sathish Kumar": "Think (L1.5)",
    "Subhani M": "Think (L1.5)",
    "Suresh R": "Think (L1.5)",
    "Vicky D": "Think (L1.5)",
    "Vimal Kumar V": "Think (L1.5)",

    "A Afrid": "A Afrid",

    "Riyaz Pasha": "Deepak Ghosh (Billable)",
    "Rahul Bhanumurthy": "Deepak Ghosh (Billable)",
    "Ramkrishnadas Mondal": "Deepak Ghosh (Billable)",
    "Ayesha Sultana": "Deepak Ghosh (Billable)",

    "Subrat Muduli": "Sreejith Menon (Support Function)",
    "Gulam Hasan": "Sreejith Menon (Support Function)",

    "Chahat Ramani": "A Afrid",
    "Faiz Ahmed": "A Afrid",
    "Steven Mario Pinto": "A Afrid",
    "Shaik Nasreen": "A Afrid",
    "Deepa S Dyamangoudar": "A Afrid",
    "Thowqeer Ahamed": "A Afrid",

    "Kanishka Kanishka": "Augustine Joseph",
    "Mohammed Roushan": "Augustine Joseph",
    "Sushmashri Sushmashri": "Augustine Joseph",
    "Anil Patode": "Augustine Joseph",

    "Suryakant Sahoo": "Avrel Fernandez",
    "Jashwanth Raj": "Avrel Fernandez",
    "Ratnashree P Mallshetty": "Avrel Fernandez",
    "Singitham Hanish": "Avrel Fernandez",
    "Dibyajyoti Panigrahi": "Avrel Fernandez",
    "Nivedita S Ganachari": "Avrel Fernandez",

    "Nikhil Kumar": "Ram Mohan Singh",
    "Mohammed Hussain": "Ram Mohan Singh",
    "Aradhya Gupta": "Ram Mohan Singh",
    "Suhas AB": "Ram Mohan Singh",

    "Vishal Pandey": "Seema Lal",
    "Pragati Priya": "Seema Lal",
    "Dhruv Mishra": "Seema Lal",
    
    "Veeresh Biradar": "Sreejith Menon",
    "Mohd Kaif": "Sreejith Menon",
    "Arvind Swami": "Sreejith Menon",

    "Annu Priya": "Syed Arbaz",
    "Gokul Anand": "Syed Arbaz",

    "Rajneesh": "Vikas B Bhovi",
    "Priyanka Hanje": "Vikas B Bhovi",
    "Swayam Panda": "Vikas B Bhovi",
    "Sakshi Bharti": "Vikas B Bhovi",

    "Subhranshu Mishra": "Sreejith Menon",
    "Pratik Kittur": "Sreejith Menon",
}


hc_lower = {
    str(k).strip().lower(): v
    for k, v in hc_dict.items()
}


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Raw Excel File (.xlsx)",
    type=["xlsx"]
)


if uploaded_file is None:

    st.info(
        "Please upload the latest raw 2-way email Excel file to generate the report."
    )

    st.stop()


# ============================================================
# READ RAW EXCEL
# ============================================================

try:

    df = pd.read_excel(
        uploaded_file,
        sheet_name=0
    )

except Exception as e:

    st.error(
        f"Unable to read the Excel file: {e}"
    )

    st.stop()


# ============================================================
# REQUIRED COLUMN VALIDATION
# ============================================================

required_columns = [
    "Current Sender (Object) (Email)",
    "Created On",
]

missing_columns = [
    column
    for column in required_columns
    if column not in df.columns
]

if missing_columns:

    st.error(
        "The uploaded Excel file is missing the following required "
        f"column(s): {', '.join(missing_columns)}"
    )

    st.stop()


# ============================================================
# AGENT / TEAM LEAD MAPPING
# ============================================================

df["Agent"] = (
    df["Current Sender (Object) (Email)"]
    .fillna("")
    .astype(str)
    .str.strip()
)

df["Mapped Team Lead"] = (
    df["Agent"]
    .str.lower()
    .map(hc_lower)
)

# Keep only mapped records for report
df_filtered = df[
    df["Mapped Team Lead"].notna()
].copy()

df_filtered["Team Lead"] = (
    df_filtered["Mapped Team Lead"]
)


# ============================================================
# DATE PREPARATION
# ============================================================

today_date = (
    pd.Timestamp.today()
    .normalize()
    .date()
)

df["Created_Date"] = pd.to_datetime(
    df["Created On"],
    errors="coerce"
).dt.date

df_filtered["Created_Date"] = pd.to_datetime(
    df_filtered["Created On"],
    errors="coerce"
).dt.date


# ============================================================
# AGEING CALCULATION
# ============================================================

def calculate_ageing(created_date, report_date):

    if pd.isna(created_date):
        return 0

    return int(
        (report_date - created_date).days
    )


df["Ageing (Days)"] = (
    df["Created_Date"]
    .apply(
        lambda x: calculate_ageing(
            x,
            today_date
        )
    )
)

df_filtered["Age"] = (
    df_filtered["Created_Date"]
    .apply(
        lambda x: calculate_ageing(
            x,
            today_date
        )
    )
)


# ============================================================
# TL SUMMARY
# ============================================================

tl_summary = (
    df_filtered
    .groupby("Team Lead")
    .size()
    .reset_index(
        name="Count of Team Lead"
    )
    .sort_values(
        "Count of Team Lead",
        ascending=False
    )
    .reset_index(drop=True)
)

grand_total = int(
    tl_summary["Count of Team Lead"].sum()
)


# ============================================================
# AGEING COLUMNS
# ============================================================

observed_ages = (
    df_filtered["Age"]
    .dropna()
    .astype(int)
    .tolist()
)

max_age = max(
    (age for age in observed_ages if age > 0),
    default=0,
)

age_columns = sorted(
    set(observed_ages)
    | set(range(1, max_age + 1))
)


# ============================================================
# BUILD TL + AGENT REPORT
# ============================================================

report_rows = []


for tl in tl_summary["Team Lead"]:

    tl_data = df_filtered[
        df_filtered["Team Lead"] == tl
    ].copy()

    tl_total = len(tl_data)

    # --------------------------------------------------------
    # TL ROW
    # --------------------------------------------------------

    tl_row = {
        "Team Lead": tl,
        "Agent": "",
        "Row Type": "TL",
    }

    for age in age_columns:

        tl_row[age] = int(
            (
                tl_data["Age"] == age
            ).sum()
        )

    tl_row["Grand Total"] = int(
        tl_total
    )

    report_rows.append(
        tl_row
    )

    # --------------------------------------------------------
    # AGENT SUMMARY
    # --------------------------------------------------------

    agent_summary = (
        tl_data
        .groupby("Agent")
        .size()
        .reset_index(
            name="Agent Total"
        )
        .sort_values(
            "Agent Total",
            ascending=False
        )
    )

    for _, agent_record in agent_summary.iterrows():

        agent = agent_record["Agent"]

        agent_data = tl_data[
            tl_data["Agent"] == agent
        ]

        agent_row = {
            "Team Lead": "",
            "Agent": agent,
            "Row Type": "Agent",
        }

        for age in age_columns:

            agent_row[age] = int(
                (
                    agent_data["Age"] == age
                ).sum()
            )

        agent_row["Grand Total"] = int(
            agent_record["Agent Total"]
        )

        report_rows.append(
            agent_row
        )


# ============================================================
# GRAND TOTAL ROW
# ============================================================

grand_row = {
    "Team Lead": "Grand Total",
    "Agent": "",
    "Row Type": "Grand Total",
}

for age in age_columns:

    grand_row[age] = int(
        (
            df_filtered["Age"] == age
        ).sum()
    )

grand_row["Grand Total"] = grand_total

report_rows.append(
    grand_row
)


tl_agent_report = pd.DataFrame(
    report_rows
)


# ============================================================
# VALIDATION
# ============================================================

tl_total_validation = int(
    tl_summary["Count of Team Lead"].sum()
)

agent_total_validation = int(
    tl_agent_report[
        tl_agent_report["Row Type"] == "Agent"
    ]["Grand Total"].sum()
)

raw_mapped_validation = len(
    df_filtered
)

if not (
    tl_total_validation
    == agent_total_validation
    == raw_mapped_validation
):

    st.error(
        "⚠️ Calculation validation failed. "
        f"TL Total={tl_total_validation}, "
        f"Agent Total={agent_total_validation}, "
        f"Mapped Raw Records={raw_mapped_validation}"
    )

    st.stop()


# ============================================================
# DISPLAY SUMMARY
# ============================================================

col1, col2, col3, col4 = st.columns(4)


with col1:

    st.markdown(
        f"""
        <div class="metric-box">
            <div class="metric-title">Report Date</div>
            <div class="metric-value">
                {today_date.strftime('%d %b %Y')}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


with col2:

    st.markdown(
        f"""
        <div class="metric-box">
            <div class="metric-title">Total Raw Records</div>
            <div class="metric-value">
                {len(df):,}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


with col3:

    st.markdown(
        f"""
        <div class="metric-box">
            <div class="metric-title">Mapped Records</div>
            <div class="metric-value">
                {grand_total:,}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


with col4:

    st.markdown(
        f"""
        <div class="metric-box">
            <div class="metric-title">Team Leads</div>
            <div class="metric-value">
                {len(tl_summary):,}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


st.success(
    f"Calculation validated successfully: "
    f"{grand_total:,} mapped records."
)


# ============================================================
# PREPARE TL & AGENT EXPORT DATA
# ============================================================

tl_agent_excel = tl_agent_report.drop(
    columns=["Row Type"]
).copy()

tl_agent_excel = tl_agent_excel[
    ["Team Lead", "Agent"]
    + age_columns
    + ["Grand Total"]
]


# ============================================================
# PREPARE TL EXPORT
# ============================================================

tl_excel = tl_summary.copy()

tl_excel = pd.concat(
    [
        tl_excel,
        pd.DataFrame(
            {
                "Team Lead": ["Grand Total"],
                "Count of Team Lead": [
                    grand_total
                ],
            }
        ),
    ],
    ignore_index=True,
)


# ============================================================
# EXCEL GENERATION
# ============================================================

excel_buffer = BytesIO()


with pd.ExcelWriter(
    excel_buffer,
    engine="openpyxl"
) as writer:

    # -----------------------------------------
    # RAW DATA
    # -----------------------------------------

    df.to_excel(
        writer,
        sheet_name="Raw Data",
        index=False
    )

    # -----------------------------------------
    # TL WISE
    # -----------------------------------------

    tl_excel.to_excel(
        writer,
        sheet_name="TL Wise",
        index=False
    )

    # -----------------------------------------
    # TL + AGENT WISE
    # -----------------------------------------

    tl_agent_excel.to_excel(
        writer,
        sheet_name="TL & Agent Wise",
        index=False
    )


excel_buffer.seek(0)


# ============================================================
# EXCEL STYLING
# ============================================================

wb = openpyxl.load_workbook(
    excel_buffer
)


# ------------------------------------------------------------
# STYLES
# ------------------------------------------------------------

header_font = Font(
    name="Arial",
    size=10,
    bold=True,
    color="FFFFFF",
)

normal_font = Font(
    name="Arial",
    size=10,
)

bold_font = Font(
    name="Arial",
    size=10,
    bold=True,
)

grand_font = Font(
    name="Arial",
    size=10,
    bold=True,
    color="FFFFFF",
)

header_fill = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid",
)

tl_fill = PatternFill(
    start_color="D9EAF7",
    end_color="D9EAF7",
    fill_type="solid",
)

grand_fill = PatternFill(
    start_color="1F4E78",
    end_color="1F4E78",
    fill_type="solid",
)

border = Border(
    left=Side(
        style="thin",
        color="B7B7B7"
    ),
    right=Side(
        style="thin",
        color="B7B7B7"
    ),
    top=Side(
        style="thin",
        color="B7B7B7"
    ),
    bottom=Side(
        style="thin",
        color="B7B7B7"
    ),
)


# ============================================================
# RAW DATA FORMATTING
# ============================================================

ws = wb["Raw Data"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False


for cell in ws[1]:

    cell.font = header_font
    cell.fill = header_fill
    cell.border = border

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


for row in ws.iter_rows(
    min_row=2,
    max_row=ws.max_row,
):

    for cell in row:

        cell.font = normal_font
        cell.border = border

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )


# ============================================================
# TL WISE FORMATTING
# ============================================================

ws = wb["TL Wise"]

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False


for cell in ws[1]:

    cell.font = header_font
    cell.fill = header_fill
    cell.border = border

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


for row_num in range(
    2,
    ws.max_row + 1
):

    for col_num in range(
        1,
        ws.max_column + 1
    ):

        cell = ws.cell(
            row=row_num,
            column=col_num
        )

        cell.font = normal_font
        cell.border = border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


# Grand total row

for col_num in range(
    1,
    ws.max_column + 1
):

    cell = ws.cell(
        row=ws.max_row,
        column=col_num
    )

    cell.font = grand_font
    cell.fill = grand_fill
    cell.border = border


# ============================================================
# TL & AGENT WISE FORMATTING
# ============================================================

ws = wb["TL & Agent Wise"]

ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False


# Header

for cell in ws[1]:

    cell.font = header_font
    cell.fill = header_fill
    cell.border = border

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


# Rows

for row_num in range(
    2,
    ws.max_row + 1
):

    team_lead = ws.cell(
        row=row_num,
        column=1
    ).value

    agent = ws.cell(
        row=row_num,
        column=2
    ).value

    # Grand Total
    if team_lead == "Grand Total":

        for col_num in range(
            1,
            ws.max_column + 1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.font = grand_font
            cell.fill = grand_fill
            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # TL subtotal
    elif team_lead and not agent:

        for col_num in range(
            1,
            ws.max_column + 1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.font = bold_font
            cell.fill = tl_fill
            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # Agent
    else:

        for col_num in range(
            1,
            ws.max_column + 1
        ):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.font = normal_font
            cell.border = border

            if col_num == 2:

                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center"
                )

            else:

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )


# ============================================================
# NUMBER FORMATTING
# ============================================================

for sheet_name in [
    "TL Wise",
    "TL & Agent Wise"
]:

    ws = wb[sheet_name]

    for row in ws.iter_rows():

        for cell in row:

            if isinstance(
                cell.value,
                (
                    int,
                    float,
                    np.integer,
                    np.floating
                )
            ):

                cell.number_format = "#,##0"


# ============================================================
# COLUMN WIDTHS
# ============================================================

for sheet_name in wb.sheetnames:

    ws = wb[sheet_name]

    for column_cells in ws.columns:

        max_length = 0

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            max_length = max(
                max_length,
                len(value)
            )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 3,
                12
            ),
            45
        )


# Special widths

wb["Raw Data"].freeze_panes = "A2"

wb["TL Wise"].column_dimensions[
    "A"
].width = 32

wb["TL Wise"].column_dimensions[
    "B"
].width = 22

wb["TL & Agent Wise"].column_dimensions[
    "A"
].width = 30

wb["TL & Agent Wise"].column_dimensions[
    "B"
].width = 32


# Age columns

for col_num in range(
    3,
    wb["TL & Agent Wise"].max_column + 1
):

    wb["TL & Agent Wise"].column_dimensions[
        get_column_letter(col_num)
    ].width = 9


# ============================================================
# FINAL EXCEL FILE
# ============================================================

styled_excel_buffer = BytesIO()

wb.save(
    styled_excel_buffer
)

styled_excel_data = (
    styled_excel_buffer.getvalue()
)


# ============================================================
# EXCEL DOWNLOAD
# ============================================================

st.markdown("---")

st.subheader(
    "📊 Excel Report"
)

st.download_button(
    label="📥 Download Formatted Excel Report",
    data=styled_excel_data,
    file_name=(
        "Lenovo_CEC_2Way_Ageing_"
        f"{today_date.strftime('%d_%b_%Y')}.xlsx"
    ),
    mime=(
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
)


# ============================================================
# EMAIL PREVIEW
# ============================================================

st.markdown("---")

show_email = st.checkbox(
    "✉️ Show Email Preview & One-Click Copy Tool",
    value=True
)


if show_email:

    # ========================================================
    # EMAIL HTML
    # ========================================================

    email_html = f"""
<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<style>

body {{
    font-family: Arial, Helvetica, sans-serif;
    color: #263445;
    font-size: 13px;
    line-height: 1.45;
    margin: 0;
    padding: 14px;
    background: #F3F6FA;
}}

.email-container {{
    max-width: 1100px;
    margin: auto;
    padding: 22px;
    background: #FFFFFF;
    border: 1px solid #DCE4EC;
    border-radius: 10px;
}}

p {{
    margin: 8px 0 12px;
}}

.email-header {{
    margin: -22px -22px 18px;
    padding: 20px 22px;
    color: #FFFFFF;
    background: #1F4E78;
    border-radius: 9px 9px 0 0;
}}

.email-header h1 {{
    margin: 0;
    color: #FFFFFF;
    font-size: 21px;
    line-height: 1.25;
}}

.email-header p {{
    margin: 5px 0 0;
    color: #E7F0F8;
    font-size: 12px;
}}

.report-info {{
    background: #F1F6FA;
    border: 1px solid #D5E1EA;
    border-left: 4px solid #2C7A9B;
    border-radius: 6px;
    padding: 11px 14px;
    margin: 14px 0 20px;
    color: #263445;
}}

.section-title {{
    color: #1F4E78;
    font-size: 14px;
    font-weight: bold;
    margin-top: 20px;
    margin-bottom: 7px;
    padding: 7px 10px;
    background: #F1F6FA;
    border-left: 3px solid #2C7A9B;
}}

.table-wrapper {{
    width: 100%;
    overflow-x: auto;
}}

table {{
    border-collapse: collapse;
    width: auto;
    max-width: none;
    table-layout: auto;
    margin: 7px 0 20px;
    font-family: Arial, Helvetica, sans-serif;
    font-size: 11px;
    line-height: 1.25;
}}

th {{
    background-color: #1F4E78;
    color: #FFFFFF;
    font-weight: bold;
    text-align: center;
    padding: 6px 7px;
    border: 1px solid #173B5A;
    vertical-align: middle;
    white-space: normal;
    overflow-wrap: anywhere;
}}

td {{
    border: 1px solid #D5DEE7;
    padding: 5px 7px;
    vertical-align: middle;
    white-space: normal;
    overflow-wrap: anywhere;
}}

.name-cell {{
    white-space: nowrap;
    overflow-wrap: normal;
    padding-left: 10px;
    padding-right: 10px;
}}

.age-header {{
    white-space: nowrap;
    overflow-wrap: normal;
    min-width: 22px;
}}

tbody tr:nth-child(even) td {{
    background-color: #F7F9FC;
}}

td.number {{
    text-align: center;
    white-space: nowrap;
}}

td.total {{
    text-align: center;
    font-weight: bold;
    white-space: nowrap;
    background-color: #EAF2F8;
}}

td.agent {{
    text-align: left;
    padding-left: 14px;
}}

tbody tr.tl-row td {{
    background-color: #E8F1F8;
    color: #173B5A;
    font-weight: bold;
    border-top: 2px solid #1F4E78;
}}

tbody tr.tl-row td.total {{
    background-color: #D7E7F2;
}}

tbody tr.grand-total td {{
    background-color: #1F4E78;
    color: #FFFFFF;
    font-weight: bold;
    border: 1px solid #173B5A;
}}

.signature {{
    margin-top: 25px;
    padding-top: 12px;
    border-top: 1px solid #D5DEE7;
}}

.small-text {{
    font-size: 12px;
    color: #666666;
}}

a {{
    color: #1F4E78;
}}

</style>

</head>

<body>

<div id="email-content">

<div class="email-container">

<div class="email-header">
<h1>2-Way Email Ageing Report</h1>
<p>Lenovo Customer Engagement Center &nbsp;|&nbsp; Daily pending email overview</p>
</div>

<p>Dear Associates,</p>

<p>
Kindly find the below pending 2-way emails as of today.
Request you to review and clear them at the earliest.
</p>

<p>Dear Tls,</p>

<p>
Kindly review the below details and support in getting these
cleared on priority.
</p>

<div class="report-info">

<strong>Report Date:</strong> {today_date.strftime('%d %B %Y')}
&nbsp;&nbsp; | &nbsp;&nbsp;
<strong>Total Pending 2-Way Emails:</strong> {grand_total:,}

</div>

<div class="section-title">
Lead Wise:
</div>

<div class="table-wrapper">

<table>

<thead>

<tr>

<th class="name-cell">Team Lead</th>

<th>Count of Team Lead</th>

</tr>

</thead>

<tbody>
"""


    # ========================================================
    # TL SUMMARY EMAIL
    # ========================================================

    for _, row in tl_summary.iterrows():

        email_html += f"""
<tr>

<td class="name-cell">
<strong>{escape(str(row["Team Lead"]))}</strong>
</td>

<td class="number">
<strong>{int(row["Count of Team Lead"]):,}</strong>
</td>

</tr>
"""


    email_html += f"""

<tr class="grand-total">

<td class="name-cell">
Grand Total
</td>

<td class="number">
{grand_total:,}
</td>

</tr>

</tbody>

</table>

</div>


<div class="section-title">
Agent Wise Count &amp; Ageing Overview:
</div>

<div class="table-wrapper">

<table>

<thead>

<tr>

<th class="name-cell">Team Lead</th>

<th class="name-cell">Agent</th>
"""


    # ========================================================
    # AGEING HEADERS
    # ========================================================

    for age in age_columns:

        email_html += f"""
<th class="age-header">{escape(str(age))}</th>
"""


    email_html += """

<th>Grand Total</th>

</tr>

</thead>

<tbody>
"""


    # ========================================================
    # TL + AGENT EMAIL TABLE
    # ========================================================

    for _, row in tl_agent_report.iterrows():

        row_type = row["Row Type"]

        # ----------------------------------------------------
        # TL ROW
        # ----------------------------------------------------

        if row_type == "TL":

            email_html += """
<tr class="tl-row">

<td class="name-cell">
"""

            email_html += escape(str(row["Team Lead"]))

            email_html += """
</td>

<td class="name-cell"></td>
"""


        # ----------------------------------------------------
        # GRAND TOTAL
        # ----------------------------------------------------

        elif row_type == "Grand Total":

            email_html += """
<tr class="grand-total">

<td class="name-cell">
Grand Total
</td>

<td class="name-cell"></td>
"""


        # ----------------------------------------------------
        # AGENT ROW
        # ----------------------------------------------------

        else:

            email_html += """
<tr>

<td class="name-cell"></td>

<td class="agent name-cell">
"""

            email_html += escape(str(row["Agent"]))

            email_html += """
</td>
"""


        # ----------------------------------------------------
        # AGEING VALUES
        # ----------------------------------------------------

        for age in age_columns:

            value = int(
                row[age]
            )

            if value == 0:

                display_value = "&nbsp;"

            else:

                display_value = f"{value:,}"


            email_html += f"""
<td class="number">
{display_value}
</td>
"""


        # ----------------------------------------------------
        # GRAND TOTAL COLUMN
        # ----------------------------------------------------

        email_html += f"""

<td class="total">
{int(row["Grand Total"]):,}
</td>

</tr>
"""


    email_html += f"""

</tbody>

</table>

</div>


<div class="signature">

<p>

Regards,<br><br>

<strong>Preeti Barik</strong><br>

Escalation SPOC – Technical Support Customer Engagement Center<br>

Lenovo India Pvt Ltd<br>

Escalation Level 2:
Sreejith Menon -
<a href="mailto:smenon4@lenovo.com">
smenon4@lenovo.com
</a>

</p>

</div>

</div>

</div>


<br>


<button

onclick="copyEmailContent()"

style="
background-color:#1F4E78;
color:#FFFFFF;
padding:11px 20px;
border:none;
border-radius:5px;
font-weight:bold;
cursor:pointer;
font-size:14px;
">

📋 Copy Formatted Email to Clipboard

</button>


<span

id="copy-status"

style="
margin-left:10px;
font-family:Arial;
font-size:13px;
color:green;
font-weight:bold;
">

</span>


<script>

function copyEmailContent() {{

    var emailDiv =
        document.getElementById("email-content");

    var range =
        document.createRange();

    range.selectNode(emailDiv);

    var selection =
        window.getSelection();

    selection.removeAllRanges();

    selection.addRange(range);

    try {{

        var successful =
            document.execCommand("copy");

        var statusMsg =
            document.getElementById("copy-status");

        if (successful) {{

            statusMsg.innerText =
                "Copied successfully! You can now paste into Outlook.";

        }} else {{

            statusMsg.innerText =
                "Copy failed.";

        }}

    }} catch (err) {{

        document.getElementById(
            "copy-status"
        ).innerText =
            "Copy failed.";

    }}

    selection.removeAllRanges();

}}

</script>

</body>

</html>
"""


    # ========================================================
    # EMAIL PREVIEW
    # ========================================================

    st.components.v1.html(
        email_html,
        height=1100,
        scrolling=True
    )


    # ========================================================
    # DOWNLOAD EMAIL HTML
    # ========================================================

    st.download_button(
        label="📥 Download Email as HTML",
        data=email_html,
        file_name=(
            "2_Way_Email_Ageing_"
            f"{today_date.strftime('%d_%b_%Y')}.html"
        ),
        mime="text/html",
    )
